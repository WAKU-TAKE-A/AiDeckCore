from pptx.util import Inches, Pt
from ..render_context import SlideContext
from ..height_estimator import ELEMENT_GAP

def render(element, ctx: SlideContext, x, y, w, h) -> float:
    rows = len(element.rows) + 1 if element.headers else len(element.rows)
    cols = len(element.headers) if element.headers else (len(element.rows[0]) if element.rows else 1)
    
    ph = ctx.find_placeholder(getattr(element, 'placeholder', None))
    target_x = ph.left if ph else x
    target_y = ph.top if ph else y
    target_w = ph.width if ph else w
    
    table_shape = ctx.slide.shapes.add_table(rows, cols, target_x, target_y, target_w, Inches(1))
    table = table_shape.table
    
    row_offset = 0
    if element.headers:
        for i, header in enumerate(element.headers):
            table.cell(0, i).text = header
        row_offset = 1
            
    for r_idx, row in enumerate(element.rows):
        for c_idx, cell in enumerate(row):
            cell_obj = table.cell(r_idx + row_offset, c_idx)
            cell_obj.text = str(cell)
            for p in cell_obj.text_frame.paragraphs:
                p.font.name = ctx.theme.font_name
                p.font.size = Pt(ctx.deck.font_size_l1) if ctx.deck.font_size_l1 else ctx.theme.size_body_small
            
    if not ph:
        rendered_height = getattr(element, 'height_hint', None)
        if rendered_height is None:
            rendered_height = table_shape.height
        return y + rendered_height + ELEMENT_GAP
    return y
